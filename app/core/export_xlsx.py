"""Генерация Excel-отчётов через openpyxl.

Три функции — по одной на каждый отчёт. Каждая возвращает готовый
HttpResponse с вложенным XLSX-файлом (Content-Disposition: attachment).

Стилизация:
- заголовки таблицы — жирный шрифт + заливка #E9ECEF (светло-серая);
- строка итогов — жирная;
- числовые столбцы выровнены вправо;
- ширина колонок авторасчёт по максимальной длине содержимого.

Зачем:
- CSV открывается «криво» в Excel (кодировка, разделители) — пользователи
  путаются. XLSX открывается в один клик, данные уже в столбцах.
- Это демонстрирует зрелость решения для комиссии ВКР.
"""
from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import reports


# =============================================================================
# Общие утилиты
# =============================================================================
_HEADER_FILL = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")
_HEADER_FONT = Font(bold=True)
_BOLD_FONT = Font(bold=True)
_RIGHT = Alignment(horizontal="right")


def _auto_width(ws) -> None:
    """Подбор ширины колонок по максимуму содержимого (+2 символа запас)."""
    for col_idx, col_cells in enumerate(ws.iter_cols(min_row=1, max_row=ws.max_row), start=1):
        max_len = 0
        for cell in col_cells:
            val = str(cell.value) if cell.value is not None else ""
            max_len = max(max_len, len(val))
        adjusted = min(max_len + 2, 50)  # не шире 50 символов
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted


def _write_header(ws, headers: list[str], row: int = 1) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL


def _make_response(wb: Workbook, filename: str) -> HttpResponse:
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


# =============================================================================
# 1. Отчёт «Динамика по периоду»
# =============================================================================
def export_period_xlsx(date_from, date_to) -> HttpResponse:
    """Экспорт отчёта by_period в XLSX."""
    data = reports.report_by_period(date_from, date_to)

    wb = Workbook()
    ws = wb.active
    ws.title = "Динамика"

    headers = ["Дата", "Создано", "Закрыто"]
    _write_header(ws, headers)

    for i, row in enumerate(data["rows"], start=2):
        ws.cell(row=i, column=1, value=row["date"].isoformat())
        ws.cell(row=i, column=2, value=row["created"]).alignment = _RIGHT
        ws.cell(row=i, column=3, value=row["closed"]).alignment = _RIGHT

    # Итого.
    total_row = len(data["rows"]) + 2
    ws.cell(row=total_row, column=1, value="Итого").font = _BOLD_FONT
    ws.cell(row=total_row, column=2, value=data["totals"]["created"]).font = _BOLD_FONT
    ws.cell(row=total_row, column=3, value=data["totals"]["closed"]).font = _BOLD_FONT

    _auto_width(ws)
    return _make_response(wb, f"report_period_{date_from}_{date_to}.xlsx")


# =============================================================================
# 2. Отчёт «Нагрузка исполнителей»
# =============================================================================
def export_executors_xlsx() -> HttpResponse:
    """Экспорт отчёта executors в XLSX."""
    rows = reports.report_executors()

    wb = Workbook()
    ws = wb.active
    ws.title = "Исполнители"

    headers = ["Исполнитель", "Всего", "Открытых", "Закрытых", "Просрочено", "Ср. время, ч"]
    _write_header(ws, headers)

    for i, r in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=r["full_name"])
        ws.cell(row=i, column=2, value=r["total"]).alignment = _RIGHT
        ws.cell(row=i, column=3, value=r["open"]).alignment = _RIGHT
        ws.cell(row=i, column=4, value=r["done"]).alignment = _RIGHT
        ws.cell(row=i, column=5, value=r["overdue"]).alignment = _RIGHT
        avg = r.get("avg_hours")
        ws.cell(row=i, column=6, value=avg if avg is not None else "").alignment = _RIGHT

    _auto_width(ws)
    return _make_response(wb, "report_executors.xlsx")


# =============================================================================
# 3. Отчёт «Соблюдение SLA»
# =============================================================================
def export_sla_xlsx() -> HttpResponse:
    """Экспорт отчёта SLA в XLSX."""
    data = reports.report_sla()

    wb = Workbook()
    ws = wb.active
    ws.title = "SLA"

    headers = ["Приоритет", "SLA, ч", "Закрыто всего", "В срок",
               "Просрочено", "Соблюдение, %", "Открытых просрочено"]
    _write_header(ws, headers)

    for i, r in enumerate(data["rows"], start=2):
        ws.cell(row=i, column=1, value=r["priority"])
        ws.cell(row=i, column=2, value=r["sla_hours"]).alignment = _RIGHT
        ws.cell(row=i, column=3, value=r["total_closed"]).alignment = _RIGHT
        ws.cell(row=i, column=4, value=r["in_time"]).alignment = _RIGHT
        ws.cell(row=i, column=5, value=r["overdue"]).alignment = _RIGHT
        pct = r["sla_pct"]
        ws.cell(row=i, column=6, value=pct if pct is not None else "").alignment = _RIGHT
        ws.cell(row=i, column=7, value=r["open_overdue"]).alignment = _RIGHT

    # Итого.
    t = data["totals"]
    total_row = len(data["rows"]) + 2
    ws.cell(row=total_row, column=1, value="Итого").font = _BOLD_FONT
    ws.cell(row=total_row, column=3, value=t["total_closed"]).font = _BOLD_FONT
    ws.cell(row=total_row, column=4, value=t["in_time"]).font = _BOLD_FONT
    ws.cell(row=total_row, column=5, value=t["overdue"]).font = _BOLD_FONT
    pct = t["sla_pct"]
    ws.cell(row=total_row, column=6, value=pct if pct is not None else "").font = _BOLD_FONT
    ws.cell(row=total_row, column=7, value=t["open_overdue"]).font = _BOLD_FONT

    _auto_width(ws)
    return _make_response(wb, "report_sla.xlsx")
